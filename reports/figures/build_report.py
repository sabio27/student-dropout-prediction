"""
Génération du rapport final — Prédiction du risque d'abandon scolaire
Auteur : Sabio
Cours  : Technologie de l'Intelligence Artificielle

Exécution : python reports/build_report.py
Génère    : reports/rapport_final.docx
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ======================================================================
#  PARAMÈTRES — modifie ici tes infos personnelles
# ======================================================================
AUTEUR = "Koffi Koffi Ambroise"          # 
WHATSAPP = "+225 05 44 80 17 44"          # 
ETABLISSEMENT = "Univ Félix Houphouët-Boigny | Univ de Rennes 2"  #
ENSEIGNANT = "Dr KAMAGATE"
DATE_REMISE = "20 avril 2026"

# ======================================================================
#  Chemins
# ======================================================================
BASE_DIR = Path(__file__).resolve().parent.parent
FIGURES_DIR = BASE_DIR / "reports" / "figures"
OUT_PATH = BASE_DIR / "reports" / "rapport_final.docx"

# ======================================================================
#  Styles et helpers
# ======================================================================
def set_run(run, bold=False, size=11, color=None, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def add_heading(doc, text, level=1, color=None):
    """Ajoute un titre avec style cohérent."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    sizes = {0: 22, 1: 16, 2: 13, 3: 11}
    run.font.size = Pt(sizes.get(level, 11))
    run.bold = True
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    else:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)  # bleu foncé
    return p

def add_paragraph(doc, text, bold=False, italic=False, size=11, align=None, justify=True):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.name = "Calibri"
        r2 = p.add_run(" " + text)
        r2.font.size = Pt(11)
        r2.font.name = "Calibri"
    else:
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.name = "Calibri"
    return p

def add_figure(doc, figure_name, caption=None, width_cm=14):
    """Insère une figure avec légende."""
    fig_path = FIGURES_DIR / figure_name
    if fig_path.exists():
        doc.add_picture(str(fig_path), width=Cm(width_cm))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(caption)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    else:
        add_paragraph(doc, f"[Figure manquante : {figure_name}]", italic=True)

def add_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Header
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = "Calibri"
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Rows
    for ri, row in enumerate(rows, start=1):
        cells = table.rows[ri].cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10)
            r.font.name = "Calibri"
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Widths
    if col_widths_cm:
        for row in table.rows:
            for cell, w in zip(row.cells, col_widths_cm):
                cell.width = Cm(w)
    return table

def add_page_break(doc):
    doc.add_page_break()

# ======================================================================
#  Création du document
# ======================================================================
doc = Document()

# Marges
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

# Style par défaut
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# ======================================================================
#  PAGE DE GARDE
# ======================================================================
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("RAPPORT DE PROJET")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prédiction du Risque")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("d'Abandon Scolaire")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Mini-projet de Machine Learning — Classification binaire")
r.italic = True
r.font.size = Pt(12)

for _ in range(6):
    doc.add_paragraph()

# Infos auteur
infos = [
    ("Auteur", AUTEUR),
    ("Contact WhatsApp", WHATSAPP),
    ("Cours", "Technologie de l'Intelligence Artificielle"),
    ("Formation", ETABLISSEMENT),
    ("Enseignant", ENSEIGNANT),
    ("Date de remise", DATE_REMISE),
]

for label, value in infos:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"{label} : ")
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)

add_page_break(doc)

# ======================================================================
#  SOMMAIRE (statique)
# ======================================================================
add_heading(doc, "Sommaire", level=1)

sommaire = [
    "1. Compréhension du problème (Business Understanding)",
    "2. Présentation du dataset",
    "3. Prétraitement des données",
    "4. Analyse exploratoire (EDA)",
    "5. Feature Engineering",
    "6. Modélisation",
    "7. Entraînement et évaluation",
    "8. Optimisation des hyperparamètres",
    "9. Déploiement (API Flask + Streamlit)",
    "10. Conclusion et perspectives",
    "Annexes",
]
for item in sommaire:
    add_paragraph(doc, item)

add_page_break(doc)

# ======================================================================
#  1. BUSINESS UNDERSTANDING
# ======================================================================
add_heading(doc, "1. Compréhension du problème", level=1)

add_heading(doc, "1.1 Contexte", level=2)
add_paragraph(doc,
    "Dans le cadre de l'enseignement supérieur, le décrochage étudiant constitue "
    "un enjeu majeur pour les établissements comme pour les apprenants. "
    "L'observation de la répartition des âges dans le dataset (15-24 ans, moyenne 19,3) "
    "situe ce travail dans un contexte d'études supérieures. L'identification précoce "
    "des étudiants en difficulté permet de mettre en place des actions correctives "
    "avant l'abandon effectif.")

add_heading(doc, "1.2 Problème à résoudre", level=2)
add_paragraph(doc,
    "À partir des caractéristiques académiques (moyenne générale, temps d'étude, "
    "taux d'absence) et socio-démographiques (âge, genre, accès Internet, activités "
    "extrascolaires) d'un étudiant, il s'agit de prédire si cet étudiant présente "
    "un risque d'abandon de ses études.")

add_heading(doc, "1.3 Type de problème Machine Learning", level=2)
add_bullet(doc, "la cible est une catégorie et non une valeur continue.", bold_prefix="Classification :")
add_bullet(doc, "deux classes possibles : 0 (pas de risque) ou 1 (risque d'abandon).", bold_prefix="Binaire :")
add_bullet(doc, "le dataset fournit la vérité terrain (dropout_risk).", bold_prefix="Supervisée :")

add_heading(doc, "1.4 Intérêts et enjeux", level=2)
add_bullet(doc, "Amélioration de la réussite étudiante.", bold_prefix="Prévention :")
add_bullet(doc, "Identifier les étudiants à risque avant l'abandon effectif.", bold_prefix="Détection précoce :")
add_bullet(doc, "Concentrer tutorat, mentorat et aides sur les profils prioritaires.", bold_prefix="Allocation des ressources :")
add_bullet(doc, "Suivi et objectivation des décisions d'accompagnement.", bold_prefix="Pilotage :")

# ======================================================================
#  2. DATASET
# ======================================================================
add_heading(doc, "2. Présentation du dataset", level=1)

add_paragraph(doc,
    "Le dataset fourni contient 300 étudiants décrits par 8 variables "
    "(7 prédicteurs + 1 cible). Aucune valeur manquante n'a été détectée, "
    "ce qui simplifie la phase de nettoyage.")

add_heading(doc, "2.1 Description des variables", level=2)
add_table(
    doc,
    headers=["Variable", "Type", "Description", "Domaine"],
    rows=[
        ["age", "int", "Âge de l'étudiant", "15-24 ans"],
        ["gender", "str", "Sexe", "Male / Female"],
        ["average_grade", "float", "Moyenne générale sur 20", "5.03-17.96"],
        ["absenteeism_rate", "float", "Taux d'absence", "0-0.5"],
        ["internet_access", "str", "Accès Internet", "Yes / No"],
        ["study_time_hours", "float", "Temps d'étude journalier (heures)", "0-5"],
        ["extra_activities", "str", "Activités extrascolaires", "Yes / No"],
        ["dropout_risk", "int", "Variable cible (risque d'abandon)", "0 / 1"],
    ],
    col_widths_cm=[3.5, 1.5, 7, 3]
)

add_heading(doc, "2.2 Distribution de la variable cible", level=2)
add_paragraph(doc,
    "La distribution des classes révèle un déséquilibre modéré : 227 étudiants "
    "(75,7%) ne présentent pas de risque d'abandon, contre 73 étudiants (24,3%) "
    "considérés à risque. Ce déséquilibre, sans être critique, impose d'utiliser "
    "des métriques robustes au-delà de la simple accuracy (F1-score, recall).")
add_figure(doc, "01_target_distribution.png",
           "Figure 1 — Distribution de la variable cible dropout_risk")

add_page_break(doc)

# ======================================================================
#  3. PRÉTRAITEMENT
# ======================================================================
add_heading(doc, "3. Prétraitement des données", level=1)

add_heading(doc, "3.1 Valeurs manquantes", level=2)
add_paragraph(doc,
    "Aucune valeur manquante n'est présente dans le dataset. Cette vérification "
    "a été effectuée via df.isnull().sum(), qui retourne 0 pour chaque variable. "
    "Aucune imputation n'est donc nécessaire.")

add_heading(doc, "3.2 Encodage des variables catégorielles", level=2)
add_paragraph(doc,
    "Les trois variables catégorielles (gender, internet_access, extra_activities) "
    "ont été encodées via One-Hot Encoding avec l'option drop_first=True afin "
    "d'éviter la multicolinéarité parfaite (piège du dummy). Cela produit trois "
    "nouvelles colonnes binaires : gender_Male, internet_access_Yes, extra_activities_Yes.")

add_heading(doc, "3.3 Normalisation", level=2)
add_paragraph(doc,
    "Une normalisation Z-score (StandardScaler de scikit-learn) a été appliquée "
    "aux variables numériques. Le scaler a été ajusté (fit) exclusivement sur "
    "l'ensemble d'entraînement, puis appliqué (transform) au test set afin "
    "d'éviter tout data leakage. Cette étape est particulièrement importante "
    "pour la Logistic Regression et le SVM, qui sont sensibles à l'échelle des données.")

# ======================================================================
#  4. EDA
# ======================================================================
add_heading(doc, "4. Analyse exploratoire (EDA)", level=1)

add_heading(doc, "4.1 Distribution des variables numériques", level=2)
add_paragraph(doc,
    "Les histogrammes par classe mettent en évidence une nette séparation "
    "sur trois des quatre variables numériques : les étudiants à risque "
    "présentent des moyennes plus faibles, des taux d'absence plus élevés "
    "et un temps d'étude journalier réduit. À l'inverse, la variable age "
    "ne semble pas discriminante.")
add_figure(doc, "02_histograms.png",
           "Figure 2 — Distribution des variables numériques par classe")

add_heading(doc, "4.2 Variables catégorielles", level=2)
add_paragraph(doc,
    "Les variables catégorielles (gender, internet_access, extra_activities) "
    "présentent des taux d'abandon très similaires entre leurs modalités "
    "(écarts < 5 points), suggérant qu'elles apportent peu d'information prédictive.")
add_figure(doc, "03_categorical.png",
           "Figure 3 — Répartition des classes selon les variables catégorielles")

add_heading(doc, "4.3 Matrice de corrélation", level=2)
add_paragraph(doc,
    "L'analyse des corrélations avec la cible confirme les observations précédentes. "
    "Les trois variables les plus corrélées à dropout_risk sont absenteeism_rate "
    "(+0.41), average_grade (-0.38) et study_time_hours (-0.30). L'âge présente "
    "une corrélation quasi nulle (-0.03).")
add_figure(doc, "04_correlation.png",
           "Figure 4 — Matrice de corrélation des variables numériques", width_cm=12)

add_page_break(doc)

# ======================================================================
#  5. FEATURE ENGINEERING
# ======================================================================
add_heading(doc, "5. Feature Engineering", level=1)

add_paragraph(doc,
    "Deux nouvelles variables ont été créées pour renforcer le pouvoir "
    "prédictif du modèle, conformément aux consignes de l'énoncé.")

add_heading(doc, "5.1 Ratio présence / absence", level=2)
add_paragraph(doc,
    "Le ratio présence/absence est défini par :")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("presence_absence_ratio = (1 - absenteeism_rate) / (absenteeism_rate + ε)")
r.italic = True
r.font.name = "Consolas"
r.font.size = Pt(10)

add_paragraph(doc,
    "Un epsilon de 1e-6 est ajouté au dénominateur pour éviter les divisions par zéro "
    "(étudiants sans absences). Le ratio est plafonné à 20 (winsorization) afin "
    "d'éviter l'introduction de valeurs aberrantes qui perturberaient la normalisation Z-score.")

add_heading(doc, "5.2 Score global pondéré", level=2)
add_paragraph(doc,
    "Un score global composite combine les trois indicateurs clés avec une "
    "pondération métier :")
add_bullet(doc, "Moyenne normalisée sur 20 (performance académique).", bold_prefix="50% — ")
add_bullet(doc, "Taux de présence (engagement).", bold_prefix="30% — ")
add_bullet(doc, "Temps d'étude normalisé sur 5h (travail personnel).", bold_prefix="20% — ")

add_heading(doc, "5.3 Justification et vérification", level=2)
add_paragraph(doc,
    "L'efficacité du feature engineering est validée a posteriori : global_score "
    "atteint une corrélation de -0.63 avec la variable cible, soit la plus élevée "
    "de toutes les variables. Elle devient également la variable la plus importante "
    "dans le modèle Random Forest (29,7% d'importance). Cette combinaison démontre "
    "la pertinence de l'approche.")

add_heading(doc, "5.4 Gestion de la multicolinéarité", level=2)
add_paragraph(doc,
    "Une variable attendance_rate (= 1 - absenteeism_rate) avait été envisagée, "
    "mais sa corrélation parfaite (-1.0) avec absenteeism_rate introduit une "
    "redondance néfaste pour les modèles linéaires (Logistic Regression, SVM). "
    "Elle a donc été supprimée au profit de presence_absence_ratio, qui apporte "
    "une information non-linéaire distincte.")

# ======================================================================
#  6. MODÉLISATION
# ======================================================================
add_heading(doc, "6. Modélisation", level=1)

add_heading(doc, "6.1 Protocole expérimental", level=2)
add_bullet(doc, "80% train / 20% test, stratifié sur la cible pour préserver les proportions de classes.",
           bold_prefix="Split :")
add_bullet(doc, "Validation croisée à 5 folds sur le set d'entraînement, scoring = F1.",
           bold_prefix="Validation :")
add_bullet(doc, "Random state = 42 pour garantir la reproductibilité.",
           bold_prefix="Graine :")

add_heading(doc, "6.2 Algorithmes testés", level=2)
add_bullet(doc,
    "modèle linéaire simple et rapide, servant de baseline. Adapté si les frontières sont linéaires.",
    bold_prefix="Logistic Regression :")
add_bullet(doc,
    "ensemble d'arbres de décision. Robuste, capable de capturer des règles de seuil et des interactions.",
    bold_prefix="Random Forest :")
add_bullet(doc,
    "capture des frontières non-linéaires via un noyau RBF. Nécessite la normalisation des features.",
    bold_prefix="Support Vector Machine (RBF) :")

add_page_break(doc)

# ======================================================================
#  7. ÉVALUATION
# ======================================================================
add_heading(doc, "7. Entraînement et évaluation", level=1)

add_heading(doc, "7.1 Métriques utilisées", level=2)
add_paragraph(doc,
    "Compte tenu du déséquilibre des classes (76/24), nous privilégions des "
    "métriques robustes plutôt que la seule accuracy :")
add_bullet(doc, "proportion de prédictions correctes (toutes classes confondues).", bold_prefix="Accuracy :")
add_bullet(doc, "parmi les étudiants prédits à risque, combien le sont réellement.", bold_prefix="Precision :")
add_bullet(doc, "parmi les étudiants réellement à risque, combien sont détectés.", bold_prefix="Recall :")
add_bullet(doc, "moyenne harmonique de precision et recall — métrique principale pour la sélection du modèle.", bold_prefix="F1-score :")

add_heading(doc, "7.2 Résultats comparatifs (avant optimisation)", level=2)
add_table(
    doc,
    headers=["Modèle", "CV F1 (mean ± std)", "Accuracy", "Precision", "Recall", "F1-score"],
    rows=[
        ["Logistic Regression", "0.753 ± 0.057", "0.867", "0.733", "0.733", "0.733"],
        ["Random Forest", "0.929 ± 0.045", "0.983", "1.000", "0.933", "0.966"],
        ["SVM (RBF)", "0.729 ± 0.018", "0.850", "0.750", "0.600", "0.667"],
    ],
    col_widths_cm=[3.5, 3, 1.8, 1.8, 1.8, 1.8]
)

add_figure(doc, "06_model_comparison.png",
           "Figure 5 — Comparaison des performances des trois modèles", width_cm=13)

add_heading(doc, "7.3 Matrices de confusion", level=2)
add_paragraph(doc,
    "Les matrices de confusion confirment la supériorité nette de Random Forest, "
    "qui ne commet pratiquement aucune erreur sur le test set.")
add_figure(doc, "05_confusion_matrices.png",
           "Figure 6 — Matrices de confusion des trois modèles", width_cm=15)

add_page_break(doc)

# ======================================================================
#  8. OPTIMISATION
# ======================================================================
add_heading(doc, "8. Optimisation des hyperparamètres", level=1)

add_paragraph(doc,
    "Un GridSearchCV à 5 folds avec scoring=F1 a été appliqué à chaque modèle. "
    "Le paramètre class_weight='balanced' a notamment été testé pour traiter "
    "le déséquilibre des classes.")

add_heading(doc, "8.1 Grilles explorées", level=2)
add_bullet(doc, "C ∈ {0.01, 0.1, 1, 10}, class_weight ∈ {None, balanced}.", bold_prefix="Logistic Regression :")
add_bullet(doc, "n_estimators ∈ {50, 100, 200}, max_depth ∈ {None, 5, 10, 20}, min_samples_split ∈ {2, 5, 10}, class_weight ∈ {None, balanced}.", bold_prefix="Random Forest :")
add_bullet(doc, "C ∈ {0.1, 1, 10}, gamma ∈ {scale, auto, 0.1}, kernel ∈ {rbf, linear}, class_weight ∈ {None, balanced}.", bold_prefix="SVM :")

add_heading(doc, "8.2 Gains après optimisation", level=2)
add_table(
    doc,
    headers=["Modèle", "F1 avant", "F1 après", "Gain"],
    rows=[
        ["Logistic Regression", "0.733", "0.722", "-0.011"],
        ["Random Forest", "0.966", "0.966", "0.000"],
        ["SVM", "0.667", "0.743", "+0.076"],
    ],
    col_widths_cm=[5, 3, 3, 3]
)

add_paragraph(doc,
    "SVM affiche le meilleur gain (+7,6 points de F1). Random Forest plafonne, "
    "déjà optimal avant GridSearch. Logistic Regression est stable, limité par "
    "son hypothèse de linéarité.")

add_heading(doc, "8.3 Modèle retenu", level=2)
add_paragraph(doc,
    "Le modèle final retenu est Random Forest avec les hyperparamètres "
    "n_estimators=200, min_samples_split=5, max_depth=None. Il atteint une "
    "accuracy de 98,3%, une precision de 100% sur la classe abandon (aucun "
    "faux positif) et un recall de 93,3% (14 vrais abandons détectés sur 15). "
    "Le F1-score final est de 0,966.")

add_figure(doc, "08_final_confusion_matrix.png",
           "Figure 7 — Matrice de confusion du modèle final (Random Forest)", width_cm=11)

add_heading(doc, "8.4 Importance des variables", level=2)
add_paragraph(doc,
    "L'analyse de l'importance des variables confirme la pertinence du feature "
    "engineering : global_score est la variable la plus prédictive (29,7%), "
    "suivie par les variables académiques. Les variables socio-démographiques "
    "contribuent de manière négligeable (<2%).")
add_figure(doc, "07_feature_importance.png",
           "Figure 8 — Importance des variables selon Random Forest", width_cm=13)

add_page_break(doc)

# ======================================================================
#  9. DÉPLOIEMENT
# ======================================================================
add_heading(doc, "9. Déploiement", level=1)

add_paragraph(doc,
    "Deux solutions de déploiement ont été développées conformément aux "
    "exigences de l'énoncé : une API Flask pour l'intégration programmatique "
    "et une interface Streamlit pour la démonstration visuelle.")

add_heading(doc, "9.1 API Flask (intégration technique)", level=2)
add_paragraph(doc,
    "L'API expose quatre endpoints REST au format JSON :")
add_bullet(doc, "page d'accueil avec documentation des endpoints.", bold_prefix="GET / :")
add_bullet(doc, "vérification de disponibilité du service.", bold_prefix="GET /health :")
add_bullet(doc, "métadonnées du modèle et métriques de performance.", bold_prefix="GET /model-info :")
add_bullet(doc, "prédiction à partir d'un payload JSON.", bold_prefix="POST /predict :")

add_paragraph(doc,
    "L'API applique le même pipeline de pré-traitement que l'entraînement "
    "(feature engineering + normalisation) pour garantir la cohérence des "
    "prédictions. Elle est lancée via la commande : python app/flask_api.py "
    "et écoute sur le port 5000.")

add_heading(doc, "9.2 Interface Streamlit (démonstration)", level=2)
add_paragraph(doc,
    "L'application Streamlit propose un formulaire de saisie avec sliders et "
    "listes déroulantes, affichant en temps réel la prédiction avec :")
add_bullet(doc, "un niveau de risque qualitatif (FAIBLE / MODÉRÉ / ÉLEVÉ).")
add_bullet(doc, "la probabilité d'abandon exacte (entre 0 et 1).")
add_bullet(doc, "les valeurs des features dérivées (global_score, ratio présence/absence).")
add_bullet(doc, "des recommandations d'action adaptées au profil.")

add_paragraph(doc,
    "L'application est déployée sur Streamlit Community Cloud, accessible "
    "via un lien public, et peut être lancée localement avec : "
    "streamlit run app/streamlit_app.py.")

add_heading(doc, "9.3 Architecture logicielle", level=2)
add_paragraph(doc,
    "Le modèle entraîné est sérialisé via joblib (fichier best_model.pkl) et "
    "contient : l'estimateur Random Forest, le StandardScaler ajusté, la "
    "liste ordonnée des features attendues, le nom du modèle et ses métriques. "
    "Cette sérialisation rend le modèle portable et réutilisable indépendamment "
    "du code d'entraînement.")

# ======================================================================
#  10. CONCLUSION
# ======================================================================
add_heading(doc, "10. Conclusion et perspectives", level=1)

add_heading(doc, "10.1 Synthèse", level=2)
add_paragraph(doc,
    "Le projet a permis de construire un modèle de classification capable "
    "de prédire le risque d'abandon scolaire avec une très haute précision "
    "(accuracy 98,3%, F1-score 0,966). Random Forest s'est imposé comme "
    "l'algorithme le plus adapté grâce à sa capacité à capturer les règles "
    "de seuil non-linéaires qui sous-tendent la génération de la variable cible.")

add_paragraph(doc,
    "Le feature engineering s'est révélé déterminant : la variable composite "
    "global_score est devenue le premier prédicteur du modèle, démontrant "
    "l'intérêt d'une combinaison pondérée des indicateurs académiques bruts.")

add_heading(doc, "10.2 Limites", level=2)
add_bullet(doc,
    "seulement 300 observations. Une validation sur un dataset réel et plus "
    "volumineux serait nécessaire avant une mise en production.",
    bold_prefix="Taille du dataset :")
add_bullet(doc,
    "la cible est générée à partir d'une règle déterministe (moyenne, absences, "
    "temps d'étude). Les performances exceptionnelles reflètent partiellement "
    "la capacité de Random Forest à ré-apprendre cette règle.",
    bold_prefix="Nature synthétique :")
add_bullet(doc,
    "le dataset ne contient pas de variables socio-économiques, familiales ou "
    "psychologiques qui jouent un rôle majeur dans l'abandon réel.",
    bold_prefix="Variables manquantes :")

add_heading(doc, "10.3 Perspectives", level=2)
add_bullet(doc, "enrichir le dataset avec des variables externes (origine socio-économique, parcours antérieur, engagement extra-académique).")
add_bullet(doc, "tester des approches par boosting (XGBoost, LightGBM) qui sont souvent encore plus performantes sur ce type de problème.")
add_bullet(doc, "implémenter une interprétabilité avancée via SHAP pour expliquer chaque prédiction individuellement aux conseillers pédagogiques.")
add_bullet(doc, "déployer l'API Flask sur un service cloud (Render, Railway) pour une accessibilité publique complète.")

# ======================================================================
#  ANNEXES
# ======================================================================
add_page_break(doc)
add_heading(doc, "Annexes", level=1)

add_heading(doc, "A. Environnement technique", level=2)
add_bullet(doc, "Python 3.13, venv")
add_bullet(doc, "pandas, numpy — manipulation des données")
add_bullet(doc, "scikit-learn — modélisation, normalisation, GridSearchCV")
add_bullet(doc, "matplotlib, seaborn — visualisations")
add_bullet(doc, "Flask — API REST")
add_bullet(doc, "Streamlit — interface web")
add_bullet(doc, "joblib — sérialisation du modèle")
add_bullet(doc, "Git + GitHub — versionnage et déploiement")

add_heading(doc, "B. Structure du projet", level=2)
p = doc.add_paragraph()
r = p.add_run(
"Techno_IA_Proj/\n"
"├── data/raw/student_dropout_dataset.csv\n"
"├── notebooks/student.ipynb\n"
"├── src/\n"
"├── models/best_model.pkl\n"
"├── app/\n"
"│   ├── flask_api.py\n"
"│   └── streamlit_app.py\n"
"├── reports/\n"
"│   ├── figures/\n"
"│   └── rapport_final.docx\n"
"├── requirements.txt\n"
"├── .gitignore\n"
"└── README.md"
)
r.font.name = "Consolas"
r.font.size = Pt(9)

add_heading(doc, "C. Commandes principales", level=2)
p = doc.add_paragraph()
r = p.add_run(
"# Créer l'environnement\n"
"python -m venv venv\n"
"source venv/Scripts/activate\n"
"pip install -r requirements.txt\n\n"
"# Lancer l'API Flask\n"
"python app/flask_api.py\n\n"
"# Lancer l'interface Streamlit\n"
"streamlit run app/streamlit_app.py"
)
r.font.name = "Consolas"
r.font.size = Pt(9)

# ======================================================================
#  Sauvegarde
# ======================================================================
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT_PATH))
print(f"OK — Rapport généré : {OUT_PATH}")
print(f"Taille : {OUT_PATH.stat().st_size / 1024:.1f} KB")